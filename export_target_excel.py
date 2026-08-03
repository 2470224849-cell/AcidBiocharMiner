#!/usr/bin/env python3
import argparse
import json
import re
import sys
from pathlib import Path
from typing import Dict, List, Tuple

from openpyxl import Workbook
from target_row_filter import filter_rows_keep_acid_pristine
from direct_sheet1_parser import normalize_and_convert_acid_conc


BIOCHAR_COLUMNS = [
    "filename",
    "sample_id",
    "biomass_source",
    "pyrolysis_temp_C",
    "hold_duration_h",
    "heating_rate_C_min",
    "acid_type",
    "acid_conc_mol_L",
    "acid_time_h",
    "acid_temp_C",
    "modification_sequence",
    "SSA_m2_g",
    "APS_nm",
    "TPV_cm3_g",
    "ash_percent",
    "C_percent",
    "O_percent",
    "N_percent",
    "H_percent",
    "pH_pzc",
]


ADS_COLUMNS = [
    "filename",
    "sample_id",
    "pollutant_name",
    "pH",
    "T_K",
    "Te_min",
    "SLR_g_L",
    "Qmax",
]


REAGENT_SAMPLE_KEYWORDS = [
    "sulfuric acid",
    "sulphuric acid",
    "hydrochloric acid",
    "nitric acid",
    "phosphoric acid",
    "acetic anhydride",
    "carbon disulfide",
    "carbon disulphide",
    "elemental sulfur",
    "elemental sulphur",
    "iron nitrate",
    "ammonium tetrathiomolybdate",
    "mercaptoethanol",
    "beta-mercaptoethanol",
]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Export two target tables into one Excel file with sheet1/sheet2."
    )
    parser.add_argument(
        "--raw-json",
        default="output/one_pdf_raw_target_rows.json",
        help="Path to raw target rows JSON from run_pdf_demo.py.",
    )
    parser.add_argument(
        "--result-json",
        default="output/one_pdf_result.json",
        help="Fallback result JSON if raw-json has no rows.",
    )
    parser.add_argument(
        "--xlsx",
        default="output/one_pdf_tables.xlsx",
        help="Output Excel file path.",
    )
    parser.add_argument(
        "--sample-filter",
        choices=["none", "acid_pristine"],
        default="acid_pristine",
        help="Filter extracted sample rows before export (default: acid_pristine).",
    )
    parser.add_argument(
        "--merge-complementary",
        action="store_true",
        default=True,
        help="Merge complementary partial rows before export (default: enabled).",
    )
    parser.add_argument(
        "--no-merge-complementary",
        dest="merge_complementary",
        action="store_false",
        help="Disable complementary-row merge.",
    )
    parser.add_argument(
        "--collapse-sheet1-by-sample",
        action="store_true",
        default=True,
        help="Force one row per (filename, sample_id) in sheet1 (default: enabled).",
    )
    parser.add_argument(
        "--no-collapse-sheet1-by-sample",
        dest="collapse_sheet1_by_sample",
        action="store_false",
        help="Disable one-row-per-sample collapse for sheet1.",
    )
    parser.add_argument(
        "--collapse-sheet2-by-sample",
        action="store_true",
        default=True,
        help=(
            "Collapse only duplicate/compatible sheet2 rows within the same "
            "(filename, sample_id, pollutant_name, pH, T_K, Te_min, SLR_g_L) "
            "condition group (default: enabled)."
        ),
    )
    parser.add_argument(
        "--no-collapse-sheet2-by-sample",
        dest="collapse_sheet2_by_sample",
        action="store_false",
        help="Disable final duplicate-condition collapse for sheet2.",
    )
    parser.add_argument(
        "--fill-missing-sample-id",
        action="store_true",
        default=True,
        help="Auto-fill missing sample_id using same-file evidence (default: enabled).",
    )
    parser.add_argument(
        "--no-fill-missing-sample-id",
        dest="fill_missing_sample_id",
        action="store_false",
        help="Disable auto-fill for missing sample_id.",
    )
    parser.add_argument(
        "--fill-unresolved-sample-id",
        action="store_true",
        default=True,
        help="Fill still-empty sample_id with UNRESOLVED_* placeholders after merge (default: enabled).",
    )
    parser.add_argument(
        "--no-fill-unresolved-sample-id",
        dest="fill_unresolved_sample_id",
        action="store_false",
        help="Keep unresolved sample_id as empty.",
    )
    return parser.parse_args()


def canonical_key(key: str) -> str:
    key = (key or "").strip().lower()
    key = re.sub(r"[\s\-]+", "_", key)
    key = re.sub(r"_+", "_", key)
    return key


def key_matches_target(key: str, target: str) -> bool:
    ck = canonical_key(key)
    ct = canonical_key(target)
    if ck == ct:
        return True
    return bool(re.fullmatch(rf"{re.escape(ct)}_\d+", ck))


def ensure_dict_rows(value) -> List[Dict]:
    if isinstance(value, dict):
        return [value]
    if not isinstance(value, list):
        return []
    return [row for row in value if isinstance(row, dict)]


def load_raw_rows(raw_path: Path) -> Tuple[List[Dict], List[Dict]]:
    if not raw_path.exists():
        return [], []
    obj = json.loads(raw_path.read_text(encoding="utf-8"))
    if isinstance(obj, dict):
        bio_rows = ensure_dict_rows(obj.get("biochar_modification", []))
        ads_rows = ensure_dict_rows(obj.get("adsorption_experiment", []))
        return bio_rows, ads_rows
    if not isinstance(obj, list):
        return [], []

    bio_rows: List[Dict] = []
    ads_rows: List[Dict] = []
    for item in obj:
        if not isinstance(item, dict):
            continue
        key = item.get("key", "")
        row = item.get("row")
        if not isinstance(row, dict):
            continue
        if key_matches_target(key, "biochar_modification"):
            bio_rows.append(row)
        elif key_matches_target(key, "adsorption_experiment"):
            ads_rows.append(row)
    return bio_rows, ads_rows


def load_rows_from_result(result_path: Path) -> Tuple[List[Dict], List[Dict]]:
    if not result_path.exists():
        return [], []
    obj = json.loads(result_path.read_text(encoding="utf-8"))
    results = obj.get("results", []) if isinstance(obj, dict) else []

    bio_rows: List[Dict] = []
    ads_rows: List[Dict] = []
    for item in results:
        if not isinstance(item, dict):
            continue
        data = item.get("data", {})
        if not isinstance(data, dict):
            continue
        for key, value in data.items():
            rows = ensure_dict_rows(value)
            if not rows:
                continue
            if key_matches_target(key, "biochar_modification"):
                bio_rows.extend(rows)
            elif key_matches_target(key, "adsorption_experiment"):
                ads_rows.extend(rows)
    return bio_rows, ads_rows


def normalize_cell(value):
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return value


def normalize_filename_value(value):
    if value is None:
        return ""
    text = str(value)
    return re.sub(r"^主文[\s_-]*原始[\s_-]*", "", text).strip()


def _dedupe_repeated_scalar(v):
    t = str(v or "").strip()
    if not t:
        return ""
    parts = t.split()
    if len(parts) >= 2 and len(set(parts)) == 1:
        return parts[0]
    m = re.fullmatch(r"([0-9]+(?:\.[0-9]+)?)\s+(?:\1(?:\s+)*)+", t)
    if m:
        return m.group(1)
    return t


def normalize_biochar_row(row: Dict) -> Dict:
    out = {
        "filename": normalize_filename_value(row.get("filename", "")),
        "sample_id": row.get("sample_id", ""),
        "biomass_source": row.get("biomass_source", ""),
        "pyrolysis_temp_C": row.get("pyrolysis_temp_C", ""),
        "hold_duration_h": row.get("hold_duration_h", ""),
        "heating_rate_C_min": row.get("heating_rate_C_min", ""),
        "acid_type": row.get("acid_type", ""),
        "acid_conc_mol_L": row.get("acid_conc_mol_L", ""),
        "acid_conc_original": row.get("acid_conc_original", ""),
        "acid_conc_mol_L_source": row.get("acid_conc_mol_L_source", ""),
        "acid_time_h": row.get("acid_time_h", ""),
        "acid_temp_C": row.get("acid_temp_C", ""),
        "modification_sequence": row.get("modification_sequence", ""),
        "SSA_m2_g": row.get("SSA_m2_g", ""),
        "APS_nm": row.get("APS_nm", ""),
        "TPV_cm3_g": row.get("TPV_cm3_g", row.get("PV_cm3_g", "")),
        "ash_percent": row.get("ash_percent", ""),
        "C_percent": row.get("C_percent", ""),
        "O_percent": row.get("O_percent", ""),
        "N_percent": row.get("N_percent", ""),
        "H_percent": row.get("H_percent", ""),
        "ratio_CN": row.get("ratio_CN", ""),
        "ratio_OC": row.get("ratio_OC", ""),
        "ratio_HC": row.get("ratio_HC", ""),
        "ratio_ON_C": row.get("ratio_ON_C", row.get("ratio_ON", "")),
        "pH_pzc": row.get("pH_pzc", ""),
    }
    # Normalize concentration into mol/L when possible and preserve provenance.
    if _is_empty(out.get("acid_conc_mol_L_source", "")):
        mol, src, original = normalize_and_convert_acid_conc(
            out.get("acid_conc_mol_L", ""), out.get("acid_type", "")
        )
        if original and _is_empty(out.get("acid_conc_original", "")):
            out["acid_conc_original"] = original
        if mol:
            out["acid_conc_mol_L"] = mol
            out["acid_conc_mol_L_source"] = src
        elif not _is_empty(out.get("acid_conc_mol_L", "")):
            out["acid_conc_mol_L_source"] = "reported_as_is"

    # Cleanup duplicated scalar artifacts like "1505.92 1505.92".
    for k in [
        "SSA_m2_g",
        "APS_nm",
        "TPV_cm3_g",
        "pyrolysis_temp_C",
        "hold_duration_h",
        "heating_rate_C_min",
        "acid_conc_mol_L",
        "acid_time_h",
        "acid_temp_C",
        "ash_percent",
        "C_percent",
        "O_percent",
        "N_percent",
        "H_percent",
        "pH_pzc",
    ]:
        out[k] = _dedupe_repeated_scalar(out.get(k, ""))
    return out


def normalize_ads_row(row: Dict) -> Dict:
    t_k, _ = _normalize_temperature_to_k(row.get("T_K", ""))
    return {
        "filename": normalize_filename_value(row.get("filename", "")),
        "sample_id": row.get("sample_id", ""),
        "pollutant_name": row.get("pollutant_name", ""),
        "pH": row.get("pH", ""),
        "T_K": t_k,
        "Te_min": row.get("Te_min", ""),
        "SLR_g_L": row.get("SLR_g_L", ""),
        "Qmax": row.get("Qmax", ""),
    }


def write_sheet(ws, columns: List[str], rows: List[Dict]) -> None:
    ws.append(columns)
    for row in rows:
        ws.append([normalize_cell(row.get(col, "")) for col in columns])


def _norm(v) -> str:
    if v is None:
        return ""
    return re.sub(r"\s+", " ", str(v).strip()).lower()


def _is_empty(v) -> bool:
    return _norm(v) == ""


def _format_number_for_cell(x: float) -> str:
    if abs(x - round(x)) < 1e-8:
        return str(int(round(x)))
    return f"{x:.2f}".rstrip("0").rstrip(".")


def _normalize_temperature_to_k(value) -> Tuple[str, bool]:
    """
    Normalize a temperature string to Kelvin only when Celsius is explicit.
    Conservative rule:
    - convert only when unit explicitly indicates Celsius (e.g., °C, ℃, celsius).
    - keep plain numeric values unchanged because unit is ambiguous.
    """
    s = str(value or "").strip()
    if not s:
        return "", False
    s2 = s.replace("℃", "°C").replace("ﹾ", "°")

    # Explicit Kelvin value: keep numeric part.
    if re.search(r"(?i)\b(?:k|kelvin)\b", s2):
        m = re.search(r"[-+]?\d+(?:\.\d+)?", s2)
        if not m:
            return s2, False
        return _format_number_for_cell(float(m.group(0))), False

    # Convert only explicit Celsius values.
    if re.search(r"(?i)(?:°\s*c|celsius|centigrade|deg(?:ree)?\s*c)", s2):
        m = re.search(r"[-+]?\d+(?:\.\d+)?", s2)
        if not m:
            return s2, False
        c = float(m.group(0))
        return _format_number_for_cell(c + 273.15), True

    return s2, False


def _is_si_filename(filename: str) -> bool:
    stem = Path(str(filename or "")).stem.lower()
    stem = re.sub(r"[\s._-]+", " ", stem).strip()
    return bool(
        re.search(
            r"\b(supplement|supplementary|supporting|suppinfo|appendix|si|mmc\d*|sm\d*|moesm\d*|esm\d*)\b",
            stem,
        )
    )


def _paper_key(filename: str) -> str:
    """
    Normalize main/SI filenames into one paper key.
    e.g., '... main.pdf' and '... mmc.docx' -> same key.
    """
    stem = Path(str(filename or "")).stem.lower()
    s = re.sub(r"[\s._-]+", " ", stem).strip()
    trailing = r"\b(main|supplement|supplementary|supporting|suppinfo|appendix|si|mmc\d*|sm\d*|moesm\d*|esm\d*)\b$"
    while True:
        ns = re.sub(rf"(?:\s+{trailing}|{trailing})", "", s).strip()
        if ns == s:
            break
        s = ns
    return _norm(s)


def _preferred_filename(rows: List[Dict]) -> str:
    names = [str(r.get("filename", "") or "").strip() for r in rows]
    names = [x for x in names if x]
    if not names:
        return ""
    for n in names:
        if not _is_si_filename(n):
            return n
    return names[0]


def _has_conflict(a: Dict, b: Dict, cols: List[str]) -> bool:
    soft_cols = {"acid_conc_original", "acid_conc_mol_L_source"}
    for c in cols:
        if c in soft_cols:
            continue
        va = _norm(a.get(c, ""))
        vb = _norm(b.get(c, ""))
        if va and vb and va != vb:
            return True
    return False


def _merge_into(dst: Dict, src: Dict, cols: List[str]) -> bool:
    return _merge_into_count(dst, src, cols) > 0


def _merge_into_count(dst: Dict, src: Dict, cols: List[str]) -> int:
    changed = 0
    for c in cols:
        if _is_empty(dst.get(c, "")) and not _is_empty(src.get(c, "")):
            dst[c] = src[c]
            changed += 1
    return changed


def _shared_equal_count(a: Dict, b: Dict, cols: List[str]) -> int:
    n = 0
    for c in cols:
        va = _norm(a.get(c, ""))
        vb = _norm(b.get(c, ""))
        if va and vb and va == vb:
            n += 1
    return n


def _norm_filename(v) -> str:
    return _paper_key(v)


def _norm_sid(v) -> str:
    return _norm(v)


def _is_unresolved_sid(v) -> bool:
    sid = str(v or "").strip().upper()
    return sid.startswith("UNRESOLVED_")


def _extract_scrh_temp_from_sid(sid: str) -> str:
    """
    Extract single target temperature from sample ids like:
    SCRH-S-20-160, SCRH-W-0.5-180, etc.
    Returns trailing temperature token (e.g., "160", "180") or "".
    """
    s = str(sid or "").strip().upper()
    m = re.fullmatch(r"SCRH-[A-Z]-[0-9]+(?:\.[0-9]+)?-([0-9]{2,4})", s)
    if not m:
        return ""
    return m.group(1)


def _looks_multi_numeric_value(v) -> bool:
    t = str(v or "").strip()
    if not t:
        return False
    nums = re.findall(r"[-+]?\d+(?:\.\d+)?", t)
    if len(nums) <= 1:
        return False
    # Typical multi-value patterns from extraction artifacts.
    return any(x in t for x in [",", ";", "/", " to ", " and ", "~", "—", "–"])


def _build_file_sid_candidates(
    bio_rows: List[Dict], ads_rows: List[Dict]
) -> Dict[str, List[str]]:
    out: Dict[str, set] = {}
    for row in [*bio_rows, *ads_rows]:
        fn = _norm_filename(row.get("filename", ""))
        sid = str(row.get("sample_id", "") or "").strip()
        if not fn or not sid:
            continue
        out.setdefault(fn, set()).add(sid)
    return {k: sorted(v) for k, v in out.items()}


def _exact_key(row: Dict, cols: List[str]) -> Tuple[str, ...]:
    return tuple(_norm(row.get(c, "")) for c in cols)


def _find_sid_by_exact_key(
    target_row: Dict, known_rows: List[Dict], cols: List[str]
) -> str:
    key = _exact_key(target_row, cols)
    if not any(key):
        return ""
    sid_set = set()
    for row in known_rows:
        sid = str(row.get("sample_id", "") or "").strip()
        if not sid:
            continue
        if _exact_key(row, cols) == key:
            sid_set.add(sid)
    return next(iter(sid_set)) if len(sid_set) == 1 else ""


def _find_sid_by_scored_vote(
    target_row: Dict,
    known_rows: List[Dict],
    cols: List[str],
    min_score: int = 2,
) -> str:
    """
    Score against known rows by counting equal non-empty condition fields.
    Return sid only when the top-vote sid is unique and confident.
    """
    votes: Dict[str, int] = {}
    best_score = 0
    for row in known_rows:
        sid = str(row.get("sample_id", "") or "").strip()
        if not sid:
            continue
        score = 0
        for c in cols:
            tv = _norm(target_row.get(c, ""))
            rv = _norm(row.get(c, ""))
            if tv and rv and tv == rv:
                score += 1
        if score <= 0:
            continue
        if score > best_score:
            best_score = score
        votes[sid] = votes.get(sid, 0) + score
    if best_score < min_score or not votes:
        return ""
    ranked = sorted(votes.items(), key=lambda x: x[1], reverse=True)
    if len(ranked) == 1:
        return ranked[0][0]
    if ranked[0][1] > ranked[1][1]:
        return ranked[0][0]
    return ""


def autofill_missing_sample_ids(
    bio_rows: List[Dict], ads_rows: List[Dict]
) -> Tuple[List[Dict], List[Dict], Dict[str, int]]:
    """
    Fill missing sample_id conservatively:
    1) same-file unique sid candidate;
    2) exact condition-key match to known rows in the same file;
    3) scored vote by condition overlap (sheet2 only, min_score=2).
    """
    file_sid = _build_file_sid_candidates(bio_rows, ads_rows)
    bio_known_by_file: Dict[str, List[Dict]] = {}
    ads_known_by_file: Dict[str, List[Dict]] = {}
    for row in bio_rows:
        fn = _norm_filename(row.get("filename", ""))
        if str(row.get("sample_id", "") or "").strip():
            bio_known_by_file.setdefault(fn, []).append(row)
    for row in ads_rows:
        fn = _norm_filename(row.get("filename", ""))
        if str(row.get("sample_id", "") or "").strip():
            ads_known_by_file.setdefault(fn, []).append(row)

    bio_fill = 0
    ads_fill = 0
    bio_exact_cols = [
        "biomass_source",
        "pyrolysis_temp_C",
        "acid_type",
        "acid_conc_mol_L",
        "acid_time_h",
        "acid_temp_C",
        "modification_sequence",
    ]
    ads_exact_cols = ["pollutant_name", "pH", "T_K", "Te_min", "SLR_g_L"]

    for row in bio_rows:
        if str(row.get("sample_id", "") or "").strip():
            continue
        fn = _norm_filename(row.get("filename", ""))
        candidates = file_sid.get(fn, [])
        sid = ""
        if len(candidates) == 1:
            sid = candidates[0]
        if not sid:
            sid = _find_sid_by_exact_key(
                row, bio_known_by_file.get(fn, []), bio_exact_cols
            )
        if sid:
            row["sample_id"] = sid
            bio_fill += 1
            bio_known_by_file.setdefault(fn, []).append(row)

    for row in ads_rows:
        if str(row.get("sample_id", "") or "").strip():
            continue
        fn = _norm_filename(row.get("filename", ""))
        candidates = file_sid.get(fn, [])
        sid = ""
        if len(candidates) == 1:
            sid = candidates[0]
        if not sid:
            sid = _find_sid_by_exact_key(
                row, ads_known_by_file.get(fn, []), ads_exact_cols
            )
        if not sid:
            sid = _find_sid_by_scored_vote(
                row, ads_known_by_file.get(fn, []), ads_exact_cols, min_score=2
            )
        if sid:
            row["sample_id"] = sid
            ads_fill += 1
            ads_known_by_file.setdefault(fn, []).append(row)

    return bio_rows, ads_rows, {"bio_filled": bio_fill, "ads_filled": ads_fill}


def absorb_unresolved_sheet1_rows(rows: List[Dict]) -> Tuple[List[Dict], Dict[str, int]]:
    """
    Merge UNRESOLVED_* rows into concrete sample rows in the same paper when
    there is one clear best match, then drop unresolved rows.
    """
    by_paper: Dict[str, List[Dict]] = {}
    for r in rows:
        by_paper.setdefault(_paper_key(r.get("filename", "")), []).append(r)

    stats = {
        "unresolved_rows": 0,
        "matched_rows": 0,
        "fields_filled_direct": 0,
        "fields_filled_paper_unique": 0,
        "rows_dropped": 0,
    }
    out: List[Dict] = []

    match_cols = [
        "biomass_source",
        "pyrolysis_temp_C",
        "hold_duration_h",
        "acid_type",
        "acid_conc_mol_L",
        "acid_time_h",
        "acid_temp_C",
        "modification_sequence",
    ]
    broad_cols = [
        "biomass_source",
        "pyrolysis_temp_C",
        "hold_duration_h",
        "heating_rate_C_min",
    ]

    for _, grp in by_paper.items():
        unresolved = [r for r in grp if _is_unresolved_sid(r.get("sample_id", ""))]
        concrete = [r for r in grp if not _is_unresolved_sid(r.get("sample_id", ""))]
        stats["unresolved_rows"] += len(unresolved)

        leftovers: List[Dict] = []
        for ur in unresolved:
            scored = []
            for cr in concrete:
                score = _shared_equal_count(ur, cr, match_cols)
                if score > 0:
                    scored.append((score, cr))
            if not scored:
                leftovers.append(ur)
                continue
            scored.sort(key=lambda x: x[0], reverse=True)
            top_score = scored[0][0]
            top = [x for x in scored if x[0] == top_score]
            if top_score < 2 or len(top) != 1:
                leftovers.append(ur)
                continue
            filled = _merge_into_count(top[0][1], ur, BIOCHAR_COLUMNS)
            if filled > 0:
                stats["matched_rows"] += 1
                stats["fields_filled_direct"] += filled
            else:
                leftovers.append(ur)

        # Conservative paper-level backfill from unresolved leftovers:
        # only when one field has a unique non-empty value in that paper.
        for c in broad_cols:
            unique_v = _unique_non_empty([x.get(c, "") for x in leftovers])
            if not unique_v:
                continue
            for cr in concrete:
                if _is_empty(cr.get(c, "")):
                    cr[c] = unique_v
                    stats["fields_filled_paper_unique"] += 1

        out.extend(concrete)
        stats["rows_dropped"] += len(unresolved)

    return out, stats


def absorb_unresolved_sheet2_rows(rows: List[Dict]) -> Tuple[List[Dict], Dict[str, int]]:
    """
    Merge UNRESOLVED_* adsorption rows into concrete sample rows in the same paper
    when there is one clear best match, then drop unresolved rows.
    """
    by_paper: Dict[str, List[Dict]] = {}
    for r in rows:
        by_paper.setdefault(_paper_key(r.get("filename", "")), []).append(r)

    stats = {
        "unresolved_rows": 0,
        "matched_rows": 0,
        "fields_filled": 0,
        "rows_dropped": 0,
    }
    out: List[Dict] = []
    match_cols = ["pollutant_name", "pH", "T_K", "Te_min", "SLR_g_L"]

    for _, grp in by_paper.items():
        unresolved = [r for r in grp if _is_unresolved_sid(r.get("sample_id", ""))]
        concrete = [r for r in grp if not _is_unresolved_sid(r.get("sample_id", ""))]
        stats["unresolved_rows"] += len(unresolved)

        for ur in unresolved:
            scored = []
            for cr in concrete:
                score = _shared_equal_count(ur, cr, match_cols)
                if score > 0:
                    scored.append((score, cr))
            if not scored:
                continue
            scored.sort(key=lambda x: x[0], reverse=True)
            top_score = scored[0][0]
            top = [x for x in scored if x[0] == top_score]
            if top_score < 2 or len(top) != 1:
                continue
            filled = _merge_into_count(top[0][1], ur, ADS_COLUMNS)
            if filled > 0:
                stats["matched_rows"] += 1
                stats["fields_filled"] += filled

        out.extend(concrete)
        stats["rows_dropped"] += len(unresolved)

    return out, stats


def merge_sheet1_rows(rows: List[Dict]) -> List[Dict]:
    """
    Merge complementary rows for the same material with conservative checks.
    """
    out: List[Dict] = []
    cue_cols = [
        "biomass_source",
        "pyrolysis_temp_C",
        "acid_type",
        "acid_conc_mol_L",
        "acid_time_h",
        "acid_temp_C",
        "modification_sequence",
    ]
    for row in rows:
        merged = False
        for ex in out:
            if _norm_filename(ex.get("filename", "")) != _norm_filename(
                row.get("filename", "")
            ):
                continue
            ex_sid = _norm(ex.get("sample_id", ""))
            rw_sid = _norm(row.get("sample_id", ""))
            if ex_sid and rw_sid and ex_sid != rw_sid:
                continue
            if _has_conflict(ex, row, BIOCHAR_COLUMNS):
                continue
            if (not ex_sid or not rw_sid) and _shared_equal_count(ex, row, cue_cols) < 2:
                continue
            _merge_into(ex, row, BIOCHAR_COLUMNS)
            if _is_si_filename(ex.get("filename", "")) and not _is_si_filename(
                row.get("filename", "")
            ):
                ex["filename"] = row.get("filename", ex.get("filename", ""))
            merged = True
            break
        if not merged:
            out.append(dict(row))
    return out


def merge_sheet2_rows(rows: List[Dict]) -> List[Dict]:
    """
    Merge complementary adsorption rows when keys/conditions do not conflict.
    """
    out: List[Dict] = []
    cue_cols = ["pH", "T_K", "Te_min", "SLR_g_L"]
    for row in rows:
        merged = False
        for ex in out:
            if _norm_filename(ex.get("filename", "")) != _norm_filename(
                row.get("filename", "")
            ):
                continue
            ep = _norm(ex.get("pollutant_name", ""))
            rp = _norm(row.get("pollutant_name", ""))
            if ep and rp and ep != rp:
                continue
            ex_sid = _norm(ex.get("sample_id", ""))
            rw_sid = _norm(row.get("sample_id", ""))
            if ex_sid and rw_sid and ex_sid != rw_sid:
                continue
            if _has_conflict(ex, row, ADS_COLUMNS):
                continue
            if (not ex_sid or not rw_sid) and _shared_equal_count(ex, row, cue_cols) < 1:
                continue
            _merge_into(ex, row, ADS_COLUMNS)
            if _is_si_filename(ex.get("filename", "")) and not _is_si_filename(
                row.get("filename", "")
            ):
                ex["filename"] = row.get("filename", ex.get("filename", ""))
            merged = True
            break
        if not merged:
            out.append(dict(row))
    return out


def _row_non_empty_count(row: Dict, cols: List[str]) -> int:
    n = 0
    for c in cols:
        if c in {"filename", "sample_id"}:
            continue
        if not _is_empty(row.get(c, "")):
            n += 1
    return n


def _row_text_weight(row: Dict, cols: List[str]) -> int:
    w = 0
    for c in cols:
        if c in {"filename", "sample_id"}:
            continue
        v = str(row.get(c, "") or "").strip()
        w += len(v)
    return w


def collapse_sheet1_one_row_per_sample(rows: List[Dict]) -> List[Dict]:
    """
    Force one row per (filename, sample_id) by choosing the most informative
    anchor row and filling empty cells from other rows in the same group.
    """
    groups: Dict[Tuple[str, str], List[Dict]] = {}
    for row in rows:
        key = (_norm_filename(row.get("filename", "")), _norm(row.get("sample_id", "")))
        groups.setdefault(key, []).append(row)

    out: List[Dict] = []
    for _, grp in groups.items():
        if len(grp) == 1:
            out.append(dict(grp[0]))
            continue

        ranked = sorted(
            grp,
            key=lambda r: (
                _row_non_empty_count(r, BIOCHAR_COLUMNS),
                _row_text_weight(r, BIOCHAR_COLUMNS),
            ),
            reverse=True,
        )
        base = dict(ranked[0])
        for src in ranked[1:]:
            _merge_into(base, src, BIOCHAR_COLUMNS)
        base["filename"] = _preferred_filename(grp)
        out.append(base)
    return out


def _resolve_pollutant_name(grp: List[Dict]) -> str:
    vals = [str(r.get("pollutant_name", "") or "").strip() for r in grp]
    vals = [v for v in vals if v]
    if not vals:
        return ""
    by_norm: Dict[str, List[str]] = {}
    for v in vals:
        by_norm.setdefault(_norm(v), []).append(v)
    best_norm = ""
    best_key = None
    for k, arr in by_norm.items():
        key = (len(arr), max(len(x) for x in arr))
        if best_key is None or key > best_key:
            best_norm = k
            best_key = key
    for v in vals:
        if _norm(v) == best_norm:
            return v
    return vals[0]


def _sheet2_sid_group_key(sid: str) -> str:
    s = _norm(sid)
    if not s:
        return ""
    s = re.sub(r"\b(sample|adsorbent|biochar|biosorbent)\b", "", s)
    s = re.sub(r"[\s_]+", "", s)
    s = re.sub(r"[^a-z0-9\-]+", "", s)
    return s


def _resolve_sheet2_sample_id(grp: List[Dict]) -> str:
    vals = [str(r.get("sample_id", "") or "").strip() for r in grp]
    vals = [v for v in vals if v]
    if not vals:
        return ""
    ranked = sorted(
        vals,
        key=lambda v: (
            0 if _sheet2_sid_group_key(v) == _norm(v).replace(" ", "") else 1,
            len(v),
            v.lower(),
        ),
    )
    return ranked[0]


def collapse_sheet2_one_row_per_sample(rows: List[Dict]) -> List[Dict]:
    """
    Keep distinct adsorption-condition rows in sheet2.

    For the same paper+sample, different pollutant/condition combinations
    should stay as separate rows. We only collapse rows that share the same
    normalized condition signature:
      (filename, sample_id, pollutant_name, pH, T_K, Te_min, SLR_g_L)

    Rows without sample_id are kept as-is.
    """
    groups: Dict[Tuple[str, str, str, str, str, str, str], List[Dict]] = {}
    passthrough: List[Dict] = []
    for row in rows:
        sid = _norm(row.get("sample_id", ""))
        if not sid:
            passthrough.append(dict(row))
            continue
        sid_key = _sheet2_sid_group_key(sid) or sid
        key = (
            _norm_filename(row.get("filename", "")),
            sid_key,
            _norm(row.get("pollutant_name", "")),
            _norm(row.get("pH", "")),
            _norm(row.get("T_K", "")),
            _norm(row.get("Te_min", "")),
            _norm(row.get("SLR_g_L", "")),
        )
        groups.setdefault(key, []).append(row)

    out: List[Dict] = []
    for _, grp in groups.items():
        if len(grp) == 1:
            out.append(dict(grp[0]))
            continue
        ranked = sorted(
            grp,
            key=lambda r: (
                _row_non_empty_count(r, ADS_COLUMNS),
                _row_text_weight(r, ADS_COLUMNS),
            ),
            reverse=True,
        )
        base = dict(ranked[0])
        for src in ranked[1:]:
            _merge_into(base, src, ADS_COLUMNS)
        base["filename"] = _preferred_filename(grp)
        base["sample_id"] = _resolve_sheet2_sample_id(grp)
        base["pollutant_name"] = _resolve_pollutant_name(grp)
        out.append(base)
    out.extend(passthrough)
    return out


def _paper_tag(filename: str) -> str:
    k = _paper_key(filename)
    tag = re.sub(r"[^a-z0-9]+", "", k)[:12]
    return tag or "paper"


def fill_unresolved_sample_ids_with_placeholder(
    rows: List[Dict], columns: List[str], prefix: str = "UNRESOLVED"
) -> Tuple[List[Dict], int]:
    """
    Replace remaining empty sample_id with deterministic placeholders.
    Placeholder is shared for rows with same paper+signature so downstream joins stay stable.
    """
    maps: Dict[str, Dict[Tuple[str, ...], str]] = {}
    counters: Dict[str, int] = {}
    filled = 0
    sig_cols = [c for c in columns if c not in {"filename", "sample_id"}]

    for row in rows:
        if not _is_empty(row.get("sample_id", "")):
            continue
        pkey = _paper_key(row.get("filename", ""))
        maps.setdefault(pkey, {})
        counters.setdefault(pkey, 0)
        sig = tuple(_norm(row.get(c, "")) for c in sig_cols)
        if sig not in maps[pkey]:
            counters[pkey] += 1
            maps[pkey][sig] = f"{prefix}_{_paper_tag(row.get('filename', ''))}_{counters[pkey]:03d}"
        row["sample_id"] = maps[pkey][sig]
        filled += 1

    return rows, filled


def _unique_non_empty(values: List[str]) -> str:
    vals = [str(v).strip() for v in values if str(v).strip()]
    uniq = []
    for v in vals:
        if v not in uniq:
            uniq.append(v)
    if len(uniq) == 1:
        return uniq[0]
    return ""


def _sid_family(sid: str) -> str:
    """
    Normalize sample_id to a base family id used for conservative inheritance.
    Examples:
    - HMB / NaMB / FeMB -> MB
    - BC-HNO3 / HNO3-BC -> BC
    - APBC / OPBC -> PBC
    """
    s = str(sid or "").strip().upper()
    s = re.sub(r"\s+", "", s)
    if not s:
        return ""

    if s in {"APBC", "OPBC"}:
        return "PBC"
    if s in {"BC", "DC", "HC", "SC", "MB", "PBC"}:
        return s

    # Token-level match first to avoid accidental suffix matches.
    tokens = [t for t in re.split(r"[-_/]+", s) if t]
    for t in tokens:
        if t in {"APBC", "OPBC"}:
            return "PBC"
        if t in {"BC", "DC", "HC", "SC", "MB", "PBC"}:
            return t

    compact = re.sub(r"[^A-Z0-9]", "", s)
    for base in ("PBC", "BC", "DC", "HC", "SC", "MB"):
        if compact.endswith(base) and compact != base:
            return base

    # Legacy lightweight prefix stripping.
    for pre in ("N", "A", "O", "H", "S", "W", "M"):
        if compact.startswith(pre) and len(compact) >= 3:
            s2 = compact[1:]
            if s2 in {"BC", "DC", "HC", "SC", "MB", "PBC"}:
                return s2
    return compact or s


def _is_base_sid_of_family(sample_id: str, family: str) -> bool:
    sid = _canon_sid(sample_id)
    fam = _canon_sid(family)
    return bool(sid and fam and sid == fam)


def _canon_sid(sid: str) -> str:
    s = str(sid or "").upper().strip()
    s = re.sub(r"\s+", "", s)
    s = re.sub(r"[^A-Z0-9\-.]+", "", s)
    return s


def _sid_tokens(sid: str) -> List[str]:
    return [t for t in re.split(r"[^a-z0-9]+", str(sid or "").lower()) if t]


def _canon_name_text(v: str) -> str:
    s = _norm(v)
    if not s:
        return ""
    # Keep this conservative: just normalize separators/case.
    return re.sub(r"[^a-z0-9]+", "", s)


def _looks_design_run_id_sid(sample_id: str) -> bool:
    """
    Experimental design / response-surface run ids such as B1, B3, A12.
    These may be meaningful for sheet2 condition rows, but are not material
    sample ids for sheet1.
    """
    s = str(sample_id or "").strip().upper()
    if not s:
        return False
    return bool(re.fullmatch(r"[A-Z]\d{1,3}", s))


def _looks_reference_like_ads_sid(sample_id: str) -> bool:
    """
    Comparative/reference sample labels from other papers, e.g.
    PF-Biochar */MB [38], Lychee seed biochar/MB [67].
    """
    s = str(sample_id or "").strip()
    if not s:
        return False
    if re.search(r"\[\d{1,3}(?:\s*,\s*\d{1,3})*\]", s):
        return True
    if "et al" in s.lower() and re.search(r"\b(19|20)\d{2}\b", s.lower()):
        return True
    return False


def _is_reagent_like_sid(sample_id: str) -> bool:
    s = str(sample_id or "").strip()
    if not s:
        return False
    sl = s.lower()
    su = s.upper()
    # Keep code-like sample ids (common in adsorption papers).
    if re.search(r"[A-Z]{1,6}\d{1,4}", su):
        return False
    if re.search(r"\b(BC|PBC|MB|HC|SC|DC|NHC|NSC|NDC|SCRH|RSB|SDB)\b", su):
        return False
    if "biochar" in sl or "hydrochar" in sl or "char" in sl:
        return False
    return any(k in sl for k in REAGENT_SAMPLE_KEYWORDS)


def _build_canonical_sid_map(rows: List[Dict]) -> Dict[str, Dict[str, str]]:
    """
    Build paper-level canonical sample id map:
    paper_key -> canon_sid -> preferred style string
    """
    by_paper: Dict[str, Dict[str, List[str]]] = {}
    for r in rows:
        sid = str(r.get("sample_id", "") or "").strip()
        if not sid or _is_reagent_like_sid(sid):
            continue
        pkey = _paper_key(r.get("filename", ""))
        if not pkey:
            continue
        c = _canon_sid(sid)
        if not c:
            continue
        by_paper.setdefault(pkey, {}).setdefault(c, []).append(sid)

    out: Dict[str, Dict[str, str]] = {}
    for pkey, cmap in by_paper.items():
        out[pkey] = {}
        for c, vals in cmap.items():
            # Prefer concise code-like style over long descriptive phrases.
            uniq: List[str] = []
            for v in vals:
                if v not in uniq:
                    uniq.append(v)
            ranked = sorted(
                uniq,
                key=lambda v: (
                    0 if re.search(r"[A-Z]{1,6}\d", v.upper()) else 1,
                    0 if ("+" in v or "-" in v) else 1,
                    len(v),
                    v.lower(),
                ),
            )
            out[pkey][c] = ranked[0]
    return out


def _build_preferred_filename_map(
    bio_rows: List[Dict], ads_rows: List[Dict]
) -> Dict[str, str]:
    by_paper: Dict[str, List[str]] = {}
    for r in [*bio_rows, *ads_rows]:
        fn = str(r.get("filename", "") or "").strip()
        if not fn:
            continue
        by_paper.setdefault(_paper_key(fn), []).append(fn)

    out: Dict[str, str] = {}
    for pkey, names in by_paper.items():
        uniq: List[str] = []
        for n in names:
            if n not in uniq:
                uniq.append(n)
        chosen = ""
        for n in uniq:
            if not _is_si_filename(n):
                chosen = n
                break
        out[pkey] = chosen or uniq[0]
    return out


def _map_sid_alias_to_canonical(sample_id: str, canon_map: Dict[str, str]) -> str:
    sid = str(sample_id or "").strip()
    if not sid or not canon_map:
        return sid
    c = _canon_sid(sid)
    if c in canon_map:
        return canon_map[c]

    # Common alias: "<base> modified with <modifier>".
    m = re.match(r"(?i)^\s*([A-Za-z0-9\-+]+)\s+modified\s+with\s+([A-Za-z0-9\-+]+)\s*$", sid)
    if m:
        b = _canon_sid(m.group(1))
        x = _canon_sid(m.group(2))
        for k, v in canon_map.items():
            if b and b in k and x and x in k:
                return v

    # Token-overlap fallback.
    sid_tok = set(_sid_tokens(sid))
    if sid_tok:
        best_v = ""
        best_score = 0
        for k, v in canon_map.items():
            vt = set(_sid_tokens(v))
            inter = len(sid_tok & vt)
            if inter > best_score and inter >= 2:
                best_score = inter
                best_v = v
        if best_v:
            return best_v
    return sid


def normalize_and_clean_sample_ids(
    bio_rows: List[Dict], ads_rows: List[Dict]
) -> Tuple[List[Dict], List[Dict], Dict[str, int]]:
    """
    1) Canonicalize sample_id aliases per paper.
    2) Drop reagent-like pseudo-sample rows that are not real material ids.
    """
    stats = {
        "bio_sid_alias_mapped": 0,
        "ads_sid_alias_mapped": 0,
        "bio_reagent_rows_dropped": 0,
        "ads_reagent_rows_dropped": 0,
        "bio_pollutant_name_rows_dropped": 0,
        "bio_design_run_rows_dropped": 0,
        "ads_reference_rows_dropped": 0,
        "filename_main_si_unified": 0,
    }

    canon_by_paper = _build_canonical_sid_map(bio_rows)
    preferred_filename_by_paper = _build_preferred_filename_map(bio_rows, ads_rows)

    def _map_rows(rows: List[Dict], sid_key: str) -> List[Dict]:
        out: List[Dict] = []
        for r in rows:
            row = dict(r)
            sid = str(row.get(sid_key, "") or "").strip()
            pkey = _paper_key(row.get("filename", ""))
            preferred_fn = preferred_filename_by_paper.get(pkey, "")
            cur_fn = str(row.get("filename", "") or "").strip()
            if preferred_fn and cur_fn and preferred_fn != cur_fn:
                row["filename"] = preferred_fn
                stats["filename_main_si_unified"] += 1
            mapped = _map_sid_alias_to_canonical(sid, canon_by_paper.get(pkey, {}))
            if sid and mapped and mapped != sid:
                row[sid_key] = mapped
                if sid_key == "sample_id":
                    if "pollutant_name" in row:
                        stats["ads_sid_alias_mapped"] += 1
                    else:
                        stats["bio_sid_alias_mapped"] += 1
            out.append(row)
        return out

    bio_rows = _map_rows(bio_rows, "sample_id")
    ads_rows = _map_rows(ads_rows, "sample_id")

    pollutant_by_paper: Dict[str, set] = {}
    for r in ads_rows:
        pkey = _paper_key(r.get("filename", ""))
        pname = _canon_name_text(r.get("pollutant_name", ""))
        if pkey and pname:
            pollutant_by_paper.setdefault(pkey, set()).add(pname)

    kept_bio: List[Dict] = []
    for r in bio_rows:
        sid = str(r.get("sample_id", "") or "").strip()
        pkey = _paper_key(r.get("filename", ""))
        if sid and _is_reagent_like_sid(sid):
            stats["bio_reagent_rows_dropped"] += 1
            continue
        if sid and _looks_design_run_id_sid(sid):
            stats["bio_design_run_rows_dropped"] += 1
            continue
        if sid and _canon_name_text(sid) in pollutant_by_paper.get(pkey, set()):
            stats["bio_pollutant_name_rows_dropped"] += 1
            continue
        kept_bio.append(r)

    kept_ads: List[Dict] = []
    for r in ads_rows:
        sid = str(r.get("sample_id", "") or "").strip()
        if sid and _is_reagent_like_sid(sid):
            stats["ads_reagent_rows_dropped"] += 1
            continue
        if sid and _looks_reference_like_ads_sid(sid):
            stats["ads_reference_rows_dropped"] += 1
            continue
        kept_ads.append(r)

    return kept_bio, kept_ads, stats


def _template_prefix_and_sid(sid: str) -> Tuple[str, str]:
    """
    Detect template/group ids such as SCRH-S-m-T / SCRH-W-m-T.
    Return (prefix, canonical_template_sid), else ("","").
    """
    cs = _canon_sid(sid)
    m = re.fullmatch(r"([A-Z0-9]+-[A-Z])-M-T", cs)
    if not m:
        return "", ""
    return m.group(1), cs


def _is_child_of_template_sid(sid: str, prefix: str, template_sid: str) -> bool:
    cs = _canon_sid(sid)
    if not cs or cs == template_sid:
        return False
    if not cs.startswith(prefix + "-"):
        return False
    tail = cs[len(prefix) + 1 :]
    # Child sample ids are concrete condition points and should contain numbers.
    return bool(re.search(r"\d", tail))


def propagate_sheet1_template_rows(rows: List[Dict]) -> Tuple[List[Dict], Dict[str, int]]:
    """
    Copy shared fields from template/group sample rows (e.g., *-m-T)
    to concrete child rows in the same paper, then drop template rows.
    """
    copy_cols = [
        "biomass_source",
        "pyrolysis_temp_C",
        "hold_duration_h",
        "heating_rate_C_min",
        "acid_type",
        "acid_conc_mol_L",
        "acid_conc_original",
        "acid_conc_mol_L_source",
        "acid_time_h",
        "acid_temp_C",
        "modification_sequence",
    ]
    by_paper: Dict[str, List[Dict]] = {}
    for r in rows:
        by_paper.setdefault(_paper_key(r.get("filename", "")), []).append(r)

    out: List[Dict] = []
    stats = {
        "template_rows_found": 0,
        "template_rows_dropped": 0,
        "child_rows_touched": 0,
        "fields_filled": 0,
    }

    for _, grp in by_paper.items():
        # prefix -> chosen template row
        template_by_prefix: Dict[str, Dict] = {}
        template_sid_by_prefix: Dict[str, str] = {}
        template_rows: List[Dict] = []

        for r in grp:
            sid = str(r.get("sample_id", "") or "").strip()
            prefix, tsid = _template_prefix_and_sid(sid)
            if not prefix:
                continue
            template_rows.append(r)
            score = sum(1 for c in copy_cols if not _is_empty(r.get(c, "")))
            prev = template_by_prefix.get(prefix)
            if prev is None:
                template_by_prefix[prefix] = r
                template_sid_by_prefix[prefix] = tsid
            else:
                prev_score = sum(1 for c in copy_cols if not _is_empty(prev.get(c, "")))
                if score > prev_score:
                    template_by_prefix[prefix] = r
                    template_sid_by_prefix[prefix] = tsid

        stats["template_rows_found"] += len(template_rows)
        touched_ids = set()
        prefix_with_children = set()

        for r in grp:
            sid = str(r.get("sample_id", "") or "").strip()
            for prefix, tmpl in template_by_prefix.items():
                tsid = template_sid_by_prefix[prefix]
                if not _is_child_of_template_sid(sid, prefix, tsid):
                    continue
                prefix_with_children.add(prefix)
                touched = False
                for c in copy_cols:
                    if _is_empty(r.get(c, "")) and not _is_empty(tmpl.get(c, "")):
                        r[c] = tmpl[c]
                        stats["fields_filled"] += 1
                        touched = True
                if touched:
                    touched_ids.add(_canon_sid(sid))

        stats["child_rows_touched"] += len(touched_ids)

        # Drop template rows only when at least one child exists in that template family.
        kept_grp: List[Dict] = []
        for r in grp:
            sid = str(r.get("sample_id", "") or "").strip()
            prefix, _ = _template_prefix_and_sid(sid)
            if prefix and prefix in prefix_with_children:
                stats["template_rows_dropped"] += 1
                continue
            kept_grp.append(r)
        out.extend(kept_grp)

    return out, stats


def fill_shared_sheet1_fields(rows: List[Dict]) -> Tuple[List[Dict], Dict[str, int]]:
    """
    Improve recall with conservative intra-paper backfill:
    - For same paper+acid_type, if acid condition field has unique value, fill missing.
    - For each paper, if pyrolysis/hold/heating/biomass has a unique dominant value,
      fill missing values for base samples (BC/PBC/MB family).
    """
    by_paper: Dict[str, List[Dict]] = {}
    for r in rows:
        by_paper.setdefault(_paper_key(r.get("filename", "")), []).append(r)

    cnt = {
        "acid_conc_fill": 0,
        "acid_time_fill": 0,
        "acid_temp_fill": 0,
        "pyrolysis_fill": 0,
        "hold_fill": 0,
        "heating_fill": 0,
        "biomass_fill": 0,
        "child_from_base_pyrolysis_fill": 0,
        "child_from_base_hold_fill": 0,
        "child_from_base_heating_fill": 0,
    }

    for _, grp in by_paper.items():
        # 1) same acid_type unique-value propagation
        by_acid: Dict[str, List[Dict]] = {}
        for r in grp:
            at = _norm(r.get("acid_type", ""))
            if at:
                by_acid.setdefault(at, []).append(r)
        for _, ag in by_acid.items():
            u_conc = _unique_non_empty([x.get("acid_conc_mol_L", "") for x in ag])
            u_time = _unique_non_empty([x.get("acid_time_h", "") for x in ag])
            u_temp = _unique_non_empty([x.get("acid_temp_C", "") for x in ag])
            for r in ag:
                if u_conc and _is_empty(r.get("acid_conc_mol_L", "")):
                    r["acid_conc_mol_L"] = u_conc
                    cnt["acid_conc_fill"] += 1
                if u_time and _is_empty(r.get("acid_time_h", "")):
                    r["acid_time_h"] = u_time
                    cnt["acid_time_fill"] += 1
                if u_temp and _is_empty(r.get("acid_temp_C", "")):
                    r["acid_temp_C"] = u_temp
                    cnt["acid_temp_fill"] += 1

        # 2) base-sample pyrolysis condition backfill from dominant paper-level values
        u_biomass = _unique_non_empty([x.get("biomass_source", "") for x in grp])
        u_pyro = _unique_non_empty([x.get("pyrolysis_temp_C", "") for x in grp])
        u_hold = _unique_non_empty([x.get("hold_duration_h", "") for x in grp])
        u_heat = _unique_non_empty([x.get("heating_rate_C_min", "") for x in grp])
        for r in grp:
            sid = r.get("sample_id", "")
            fam = _sid_family(sid)
            is_base = fam in {"BC", "PBC", "MB", "SC", "DC", "HC"} and _is_base_sid_of_family(
                sid, fam
            )
            if not is_base:
                continue
            if u_biomass and _is_empty(r.get("biomass_source", "")):
                r["biomass_source"] = u_biomass
                cnt["biomass_fill"] += 1
            if u_pyro and _is_empty(r.get("pyrolysis_temp_C", "")):
                r["pyrolysis_temp_C"] = u_pyro
                cnt["pyrolysis_fill"] += 1
            if u_hold and _is_empty(r.get("hold_duration_h", "")):
                r["hold_duration_h"] = u_hold
                cnt["hold_fill"] += 1
            if u_heat and _is_empty(r.get("heating_rate_C_min", "")):
                r["heating_rate_C_min"] = u_heat
                cnt["heating_fill"] += 1

        # 3) inherit thermal parameters from base sample to sibling modified samples.
        # Only fill blanks; never overwrite non-empty extracted values.
        by_family: Dict[str, List[Dict]] = {}
        for r in grp:
            fam = _sid_family(r.get("sample_id", ""))
            if fam in {"BC", "PBC", "MB", "SC", "DC", "HC"}:
                by_family.setdefault(fam, []).append(r)

        for fam, fg in by_family.items():
            base_rows = [
                x for x in fg if _is_base_sid_of_family(x.get("sample_id", ""), fam)
            ]
            if not base_rows:
                continue
            base_pyro = _unique_non_empty([x.get("pyrolysis_temp_C", "") for x in base_rows])
            base_hold = _unique_non_empty([x.get("hold_duration_h", "") for x in base_rows])
            base_heat = _unique_non_empty([x.get("heating_rate_C_min", "") for x in base_rows])
            if not (base_pyro or base_hold or base_heat):
                continue
            for r in fg:
                if _is_base_sid_of_family(r.get("sample_id", ""), fam):
                    continue
                if base_pyro and _is_empty(r.get("pyrolysis_temp_C", "")):
                    r["pyrolysis_temp_C"] = base_pyro
                    cnt["child_from_base_pyrolysis_fill"] += 1
                if base_hold and _is_empty(r.get("hold_duration_h", "")):
                    r["hold_duration_h"] = base_hold
                    cnt["child_from_base_hold_fill"] += 1
                if base_heat and _is_empty(r.get("heating_rate_C_min", "")):
                    r["heating_rate_C_min"] = base_heat
                    cnt["child_from_base_heating_fill"] += 1

    return rows, cnt


def clear_acid_fields_for_base_pristine_rows(rows: List[Dict]) -> Tuple[List[Dict], Dict[str, int]]:
    """
    Guardrail for base control samples:
    - For explicit base IDs (BC/MB/HC/SC/DC/PBC), when the same paper has sibling
      modified samples in the same family, treat the base row as pristine control.
    - Clear acid fields on such base rows unless the row explicitly states acid route
      in modification_sequence.
    This avoids accidental acid-condition leakage into base control rows.
    """
    by_paper: Dict[str, List[Dict]] = {}
    for r in rows:
        by_paper.setdefault(_paper_key(r.get("filename", "")), []).append(r)

    base_ids = {"BC", "MB", "HC", "SC", "DC", "PBC"}
    acid_fields = ("acid_type", "acid_conc_mol_L", "acid_time_h", "acid_temp_C")
    stats = {"rows_touched": 0, "fields_cleared": 0}

    for _, grp in by_paper.items():
        by_family: Dict[str, List[Dict]] = {}
        for r in grp:
            fam = _sid_family(r.get("sample_id", ""))
            if fam in base_ids:
                by_family.setdefault(fam, []).append(r)

        for fam, fg in by_family.items():
            # Sibling modified sample exists in this family.
            has_modified_sibling = any(
                not _is_base_sid_of_family(x.get("sample_id", ""), fam) for x in fg
            )
            if not has_modified_sibling:
                continue

            for r in fg:
                sid = str(r.get("sample_id", "") or "").strip()
                if not _is_base_sid_of_family(sid, fam):
                    continue
                mod_seq = _norm(r.get("modification_sequence", ""))
                # If base row explicitly says acid route, keep as-is.
                if "acid" in mod_seq:
                    continue
                cleared = 0
                for f in acid_fields:
                    if not _is_empty(r.get(f, "")):
                        r[f] = ""
                        cleared += 1
                if cleared > 0:
                    stats["rows_touched"] += 1
                    stats["fields_cleared"] += cleared

    return rows, stats


def fill_shared_sheet2_fields(rows: List[Dict]) -> Tuple[List[Dict], Dict[str, int]]:
    """
    Improve sheet2 completeness with conservative same-sample propagation.
    - Normalize explicit Celsius temperatures in T_K to Kelvin.
    - For same paper+sample+pollutant, propagate unique pH/T_K/Te_min/SLR to blanks.
    - Then for same paper+sample (across pollutant groups), propagate only when unique.
    """
    cnt = {
        "temp_c_to_k": 0,
        "group_fill_pH": 0,
        "group_fill_T_K": 0,
        "group_fill_Te_min": 0,
        "group_fill_SLR_g_L": 0,
        "sample_fill_pH": 0,
        "sample_fill_T_K": 0,
        "sample_fill_Te_min": 0,
        "sample_fill_SLR_g_L": 0,
    }

    cond_cols = ["pH", "T_K", "Te_min", "SLR_g_L"]

    # 0) normalize explicit Celsius values into Kelvin in-place.
    for r in rows:
        v, converted = _normalize_temperature_to_k(r.get("T_K", ""))
        if converted:
            cnt["temp_c_to_k"] += 1
        if v != r.get("T_K", ""):
            r["T_K"] = v

    # 1) same paper+sample+pollutant propagation.
    by_group: Dict[Tuple[str, str, str], List[Dict]] = {}
    for r in rows:
        sid = _norm(r.get("sample_id", ""))
        if not sid:
            continue
        key = (
            _paper_key(r.get("filename", "")),
            sid,
            _norm(r.get("pollutant_name", "")),
        )
        by_group.setdefault(key, []).append(r)

    for _, grp in by_group.items():
        uniq = {c: _unique_non_empty([x.get(c, "") for x in grp]) for c in cond_cols}
        for r in grp:
            for c in cond_cols:
                if uniq[c] and _is_empty(r.get(c, "")):
                    r[c] = uniq[c]
                    cnt[f"group_fill_{c}"] += 1

    # 2) same paper+sample propagation (only if still unique at sample level).
    by_sample: Dict[Tuple[str, str], List[Dict]] = {}
    for r in rows:
        sid = _norm(r.get("sample_id", ""))
        if not sid:
            continue
        key = (_paper_key(r.get("filename", "")), sid)
        by_sample.setdefault(key, []).append(r)

    for _, grp in by_sample.items():
        uniq = {c: _unique_non_empty([x.get(c, "") for x in grp]) for c in cond_cols}
        for r in grp:
            for c in cond_cols:
                if uniq[c] and _is_empty(r.get(c, "")):
                    r[c] = uniq[c]
                    cnt[f"sample_fill_{c}"] += 1

    return rows, cnt


def enforce_scrh_single_pyrolysis_temp(rows: List[Dict]) -> Tuple[List[Dict], Dict[str, int]]:
    """
    Ensure SCRH-* samples have one pyrolysis temperature per sample row.
    Rule: use trailing numeric token in sample_id (e.g., *-160 -> 160).
    Only applies when current pyrolysis_temp_C is empty or a multi-value string.
    """
    stats = {"rows_touched": 0, "rows_filled": 0, "rows_overwritten": 0}
    for r in rows:
        sid = str(r.get("sample_id", "") or "").strip()
        t = _extract_scrh_temp_from_sid(sid)
        if not t:
            continue
        stats["rows_touched"] += 1
        cur = r.get("pyrolysis_temp_C", "")
        if _is_empty(cur):
            r["pyrolysis_temp_C"] = t
            stats["rows_filled"] += 1
            continue
        if _looks_multi_numeric_value(cur):
            if str(cur).strip() != t:
                r["pyrolysis_temp_C"] = t
                stats["rows_overwritten"] += 1
    return rows, stats


def main() -> int:
    args = parse_args()
    raw_path = Path(args.raw_json).expanduser()
    result_path = Path(args.result_json).expanduser()
    xlsx_path = Path(args.xlsx).expanduser()

    bio_rows, ads_rows = load_raw_rows(raw_path)
    source = "raw-json"
    if not bio_rows and not ads_rows:
        bio_rows, ads_rows = load_rows_from_result(result_path)
        source = "result-json"

    if not bio_rows and not ads_rows:
        print(
            "[ERROR] No target rows found. Checked raw-json and result-json.",
            file=sys.stderr,
        )
        return 1

    bio_rows = [normalize_biochar_row(row) for row in bio_rows]
    ads_rows = [normalize_ads_row(row) for row in ads_rows]

    sid_norm_stats = {}
    bio_rows, ads_rows, sid_norm_stats = normalize_and_clean_sample_ids(
        bio_rows, ads_rows
    )

    template_stats = {}
    bio_rows, template_stats = propagate_sheet1_template_rows(bio_rows)

    filter_stats = {}
    if args.sample_filter == "acid_pristine":
        bio_rows, ads_rows, filter_stats = filter_rows_keep_acid_pristine(
            bio_rows, ads_rows
        )

    sid_fill_stats = {}
    before_sid_blank_bio = sum(
        1 for r in bio_rows if _is_empty(r.get("sample_id", ""))
    )
    before_sid_blank_ads = sum(
        1 for r in ads_rows if _is_empty(r.get("sample_id", ""))
    )
    if args.fill_missing_sample_id:
        bio_rows, ads_rows, sid_fill_stats = autofill_missing_sample_ids(
            bio_rows, ads_rows
        )
    sheet1_shared_fill_stats = {}
    bio_rows, sheet1_shared_fill_stats = fill_shared_sheet1_fields(bio_rows)
    unresolved_absorb_pre_sheet1_stats = {}
    unresolved_absorb_pre_sheet2_stats = {}
    bio_rows, unresolved_absorb_pre_sheet1_stats = absorb_unresolved_sheet1_rows(
        bio_rows
    )
    ads_rows, unresolved_absorb_pre_sheet2_stats = absorb_unresolved_sheet2_rows(
        ads_rows
    )
    after_sid_blank_bio = sum(
        1 for r in bio_rows if _is_empty(r.get("sample_id", ""))
    )
    after_sid_blank_ads = sum(
        1 for r in ads_rows if _is_empty(r.get("sample_id", ""))
    )

    before_merge_bio = len(bio_rows)
    before_merge_ads = len(ads_rows)
    after_merge_bio = len(bio_rows)
    after_merge_ads = len(ads_rows)
    before_collapse_bio = len(bio_rows)
    before_collapse_ads = len(ads_rows)
    sheet2_shared_fill_stats = {}
    if args.merge_complementary:
        bio_rows = merge_sheet1_rows(bio_rows)
        ads_rows = merge_sheet2_rows(ads_rows)
        after_merge_bio = len(bio_rows)
        after_merge_ads = len(ads_rows)
    ads_rows, sheet2_shared_fill_stats = fill_shared_sheet2_fields(ads_rows)
    if args.collapse_sheet1_by_sample:
        before_collapse_bio = len(bio_rows)
        bio_rows = collapse_sheet1_one_row_per_sample(bio_rows)
    base_pristine_guard_stats = {}
    bio_rows, base_pristine_guard_stats = clear_acid_fields_for_base_pristine_rows(
        bio_rows
    )

    unresolved_fill_bio = 0
    unresolved_fill_ads = 0
    if args.fill_unresolved_sample_id:
        bio_rows, unresolved_fill_bio = fill_unresolved_sample_ids_with_placeholder(
            bio_rows, BIOCHAR_COLUMNS
        )
        ads_rows, unresolved_fill_ads = fill_unresolved_sample_ids_with_placeholder(
            ads_rows, ADS_COLUMNS
        )

    # Final pass: absorb any UNRESOLVED_* placeholder rows into same-paper samples
    # when possible, and drop unresolved rows from final export.
    unresolved_absorb_post_sheet1_stats = {}
    unresolved_absorb_post_sheet2_stats = {}
    bio_rows, unresolved_absorb_post_sheet1_stats = absorb_unresolved_sheet1_rows(
        bio_rows
    )
    ads_rows, unresolved_absorb_post_sheet2_stats = absorb_unresolved_sheet2_rows(
        ads_rows
    )
    if args.collapse_sheet2_by_sample:
        before_collapse_ads = len(ads_rows)
        ads_rows = collapse_sheet2_one_row_per_sample(ads_rows)
    scrh_pyro_stats = {}
    bio_rows, scrh_pyro_stats = enforce_scrh_single_pyrolysis_temp(bio_rows)

    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "sheet1"
    ws2 = wb.create_sheet("sheet2")

    write_sheet(ws1, BIOCHAR_COLUMNS, bio_rows)
    write_sheet(ws2, ADS_COLUMNS, ads_rows)

    wb.save(xlsx_path)

    print("=== Excel Export Summary ===")
    print(f"Source: {source}")
    print(f"Sample filter: {args.sample_filter}")
    print(f"Auto-fill missing sample_id: {args.fill_missing_sample_id}")
    print(f"Fill unresolved sample_id with placeholder: {args.fill_unresolved_sample_id}")
    print(f"Merge complementary rows: {args.merge_complementary}")
    print(f"Collapse sheet1 one-row-per-sample: {args.collapse_sheet1_by_sample}")
    print(f"Collapse sheet2 one-row-per-sample: {args.collapse_sheet2_by_sample}")
    print(
        "Template propagation stats: "
        f"templates found {template_stats.get('template_rows_found', 0)}, "
        f"templates dropped {template_stats.get('template_rows_dropped', 0)}, "
        f"children touched {template_stats.get('child_rows_touched', 0)}, "
        f"fields filled {template_stats.get('fields_filled', 0)}"
    )
    print(
        "Sample-id normalization stats: "
        f"bio alias mapped +{sid_norm_stats.get('bio_sid_alias_mapped', 0)}, "
        f"ads alias mapped +{sid_norm_stats.get('ads_sid_alias_mapped', 0)}, "
        f"bio reagent rows dropped {sid_norm_stats.get('bio_reagent_rows_dropped', 0)}, "
        f"ads reagent rows dropped {sid_norm_stats.get('ads_reagent_rows_dropped', 0)}, "
        f"bio pollutant-name rows dropped {sid_norm_stats.get('bio_pollutant_name_rows_dropped', 0)}, "
        f"bio design-run rows dropped {sid_norm_stats.get('bio_design_run_rows_dropped', 0)}, "
        f"ads reference rows dropped {sid_norm_stats.get('ads_reference_rows_dropped', 0)}, "
        f"filename main/SI unified {sid_norm_stats.get('filename_main_si_unified', 0)}"
    )
    if filter_stats:
        print(
            "Filter stats: "
            f"bio {filter_stats['bio_before']} -> {filter_stats['bio_after']}, "
            f"ads {filter_stats['ads_before']} -> {filter_stats['ads_after']}"
        )
    if args.merge_complementary:
        print(
            "Merge stats: "
            f"bio {before_merge_bio} -> {after_merge_bio}, "
            f"ads {before_merge_ads} -> {after_merge_ads}"
        )
    if args.fill_missing_sample_id:
        print(
            "Sample-id fill stats: "
            f"bio +{sid_fill_stats.get('bio_filled', 0)}, "
            f"ads +{sid_fill_stats.get('ads_filled', 0)}"
        )
    print(
        "Sheet1 shared-fill stats: "
        f"acid_conc +{sheet1_shared_fill_stats.get('acid_conc_fill', 0)}, "
        f"acid_time +{sheet1_shared_fill_stats.get('acid_time_fill', 0)}, "
        f"acid_temp +{sheet1_shared_fill_stats.get('acid_temp_fill', 0)}, "
        f"pyro +{sheet1_shared_fill_stats.get('pyrolysis_fill', 0)}, "
        f"hold +{sheet1_shared_fill_stats.get('hold_fill', 0)}, "
        f"heating +{sheet1_shared_fill_stats.get('heating_fill', 0)}, "
        f"biomass +{sheet1_shared_fill_stats.get('biomass_fill', 0)}, "
        f"child_from_base_pyro +{sheet1_shared_fill_stats.get('child_from_base_pyrolysis_fill', 0)}, "
        f"child_from_base_hold +{sheet1_shared_fill_stats.get('child_from_base_hold_fill', 0)}, "
        f"child_from_base_heating +{sheet1_shared_fill_stats.get('child_from_base_heating_fill', 0)}"
    )
    print(
        "Sheet1 base-pristine guard: "
        f"rows_touched +{base_pristine_guard_stats.get('rows_touched', 0)}, "
        f"acid_fields_cleared +{base_pristine_guard_stats.get('fields_cleared', 0)}"
    )
    print(
        "Sheet2 shared-fill stats: "
        f"C_to_K +{sheet2_shared_fill_stats.get('temp_c_to_k', 0)}, "
        f"group pH +{sheet2_shared_fill_stats.get('group_fill_pH', 0)}, "
        f"group T_K +{sheet2_shared_fill_stats.get('group_fill_T_K', 0)}, "
        f"group Te_min +{sheet2_shared_fill_stats.get('group_fill_Te_min', 0)}, "
        f"group SLR +{sheet2_shared_fill_stats.get('group_fill_SLR_g_L', 0)}, "
        f"sample pH +{sheet2_shared_fill_stats.get('sample_fill_pH', 0)}, "
        f"sample T_K +{sheet2_shared_fill_stats.get('sample_fill_T_K', 0)}, "
        f"sample Te_min +{sheet2_shared_fill_stats.get('sample_fill_Te_min', 0)}, "
        f"sample SLR +{sheet2_shared_fill_stats.get('sample_fill_SLR_g_L', 0)}"
    )
    print(
        "Unresolved absorb PRE stats (sheet1): "
        f"unresolved {unresolved_absorb_pre_sheet1_stats.get('unresolved_rows', 0)}, "
        f"matched {unresolved_absorb_pre_sheet1_stats.get('matched_rows', 0)}, "
        f"direct fields +{unresolved_absorb_pre_sheet1_stats.get('fields_filled_direct', 0)}, "
        f"paper-unique fields +{unresolved_absorb_pre_sheet1_stats.get('fields_filled_paper_unique', 0)}, "
        f"dropped {unresolved_absorb_pre_sheet1_stats.get('rows_dropped', 0)}"
    )
    print(
        "Unresolved absorb PRE stats (sheet2): "
        f"unresolved {unresolved_absorb_pre_sheet2_stats.get('unresolved_rows', 0)}, "
        f"matched {unresolved_absorb_pre_sheet2_stats.get('matched_rows', 0)}, "
        f"fields +{unresolved_absorb_pre_sheet2_stats.get('fields_filled', 0)}, "
        f"dropped {unresolved_absorb_pre_sheet2_stats.get('rows_dropped', 0)}"
    )
    print(
        "Unresolved absorb POST stats (sheet1): "
        f"unresolved {unresolved_absorb_post_sheet1_stats.get('unresolved_rows', 0)}, "
        f"matched {unresolved_absorb_post_sheet1_stats.get('matched_rows', 0)}, "
        f"direct fields +{unresolved_absorb_post_sheet1_stats.get('fields_filled_direct', 0)}, "
        f"paper-unique fields +{unresolved_absorb_post_sheet1_stats.get('fields_filled_paper_unique', 0)}, "
        f"dropped {unresolved_absorb_post_sheet1_stats.get('rows_dropped', 0)}"
    )
    print(
        "Unresolved absorb POST stats (sheet2): "
        f"unresolved {unresolved_absorb_post_sheet2_stats.get('unresolved_rows', 0)}, "
        f"matched {unresolved_absorb_post_sheet2_stats.get('matched_rows', 0)}, "
        f"fields +{unresolved_absorb_post_sheet2_stats.get('fields_filled', 0)}, "
        f"dropped {unresolved_absorb_post_sheet2_stats.get('rows_dropped', 0)}"
    )
    print(
        "SCRH single-pyrolysis enforcement: "
        f"touched {scrh_pyro_stats.get('rows_touched', 0)}, "
        f"filled +{scrh_pyro_stats.get('rows_filled', 0)}, "
        f"overwritten +{scrh_pyro_stats.get('rows_overwritten', 0)}"
    )
    if args.fill_unresolved_sample_id:
        print(
            "Sample-id unresolved-placeholder stats: "
            f"bio +{unresolved_fill_bio}, ads +{unresolved_fill_ads}"
        )
    print(
        "Sample-id blanks: "
        f"bio {before_sid_blank_bio} -> {after_sid_blank_bio}, "
        f"ads {before_sid_blank_ads} -> {after_sid_blank_ads}"
    )
    if args.collapse_sheet1_by_sample:
        print(f"Collapse stats: bio {before_collapse_bio} -> {len(bio_rows)}")
    if args.collapse_sheet2_by_sample:
        print(f"Collapse stats: ads {before_collapse_ads} -> {len(ads_rows)}")
    print(f"Biochar-modification rows (sheet1): {len(bio_rows)}")
    print(f"Adsorption-experiment rows (sheet2): {len(ads_rows)}")
    print(f"Saved: {xlsx_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
